"""Create the Word reference (style) document used by pandoc: Elsevier-style review manuscript layout
(A4, 2.5 cm margins, Times New Roman 12 pt, double spacing, continuous line numbers, page numbers)."""
import subprocess, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

out = sys.argv[1]
subprocess.run(["pandoc", "-o", out, "--print-default-data-file", "reference.docx"], check=True)
with open(out, "wb") as f:
    f.write(subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], check=True, capture_output=True).stdout)
d = Document(out)
for s in d.styles:
    try:
        if s.font is not None:
            s.font.name = "Times New Roman"; s.font.color.rgb = RGBColor(0, 0, 0)
            rpr = s.element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), "Times New Roman")
    except Exception:
        pass
_by = {st.name.lower(): st for st in d.styles}
def para(name, size, bold=None, italic=None, before=0, after=6, spacing=2.0, align=None):
    s = _by.get(name.lower())
    if s is None:
        raise KeyError(name)
    s.font.size = Pt(size)
    if bold is not None: s.font.bold = bold
    if italic is not None: s.font.italic = italic
    pf = s.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = spacing
    if align is not None: pf.alignment = align
for n in ("Normal", "Body Text", "First Paragraph"):
    try: para(n, 12)
    except KeyError: pass
para("Title", 16, bold=True, after=12, spacing=1.5) if "title" in _by else None
for n, sz in (("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 12)):
    if n.lower() in _by: para(n, sz, bold=True, italic=(n == "Heading 3"), before=12, after=6, spacing=1.5)
for n in ("Caption", "Image Caption", "Table Caption"):
    try: para(n, 10, italic=False, after=10, spacing=1.15)
    except KeyError: pass
para("Compact", 9.5, after=0, spacing=1.0)
try: para("Abstract", 12)
except KeyError: pass
sec = d.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))
ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "283")
sec._sectPr.append(ln)
# page number footer
p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
p.alignment = 1
r = p.add_run()
for t, v in (("begin", None), (None, "PAGE"), ("end", None)):
    if t:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._r.append(e)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = v; r._r.append(e)
d.save(out)
print("reference docx written", out)
