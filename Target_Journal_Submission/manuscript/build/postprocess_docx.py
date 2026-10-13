"""Post-process the pandoc DOCX: three-line (booktabs-like) tables, centred figures,
page break before the reference list, keep captions with their objects."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def _border(el, sides, sz=8):
    tcPr = el.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for side in sides:
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), "000000")
        b.append(e)

def postprocess(path):
    d = Document(path)
    for t in d.tables:
        t.alignment = 1
        rows = t.rows
        for c in rows[0].cells:
            _border(c._tc, ("top", "bottom"))
        for c in rows[-1].cells:
            _border(c._tc, ("bottom",))
        for r in rows:
            for c in r.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = 0
    for p in d.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
        if p.style.name in ("Table Caption",):
            p.paragraph_format.keep_with_next = True
        if p.style.name.startswith("Heading") and p.text.strip() == "References":
            p.insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)
    d.save(path)
